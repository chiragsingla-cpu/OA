"""Parse → chunk → embed → store."""

import io
from pathlib import Path

from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PasswordType, PdfReader

from app import vectorstore
from app.config import settings
from app.embeddings import embed_passages

TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}


class UnsupportedFileType(ValueError):
    pass


class EmptyDocument(ValueError):
    pass


class UnreadableDocument(ValueError):
    """The file could not be parsed (damaged or password-protected). The message is shown to the admin."""


def _pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        # PDFs locked only against editing open with an empty password; ones that need a password to open don't.
        if reader.is_encrypted and reader.decrypt("") == PasswordType.NOT_DECRYPTED:
            raise UnreadableDocument("This PDF is password-protected. Remove the password and upload it again.")
        return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()
    except UnreadableDocument:
        raise
    except Exception as error:  # pypdf raises many error types for broken files
        raise UnreadableDocument(
            "Could not read this PDF. It may be damaged; try saving or exporting it again."
        ) from error


def _docx_text(data: bytes) -> str:
    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as error:  # not a valid .docx package (damaged, or an old .doc renamed)
        raise UnreadableDocument(
            "Could not read this Word file. It may be damaged; save it again as .docx and re-upload."
        ) from error
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()


def extract_text(filename: str, data: bytes) -> str:
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        return _pdf_text(data)

    if extension == ".docx":
        return _docx_text(data)

    if extension in TEXT_EXTENSIONS:
        return data.decode("utf-8", errors="replace").strip()

    raise UnsupportedFileType(f"Unsupported file type '{extension}'. Use PDF, DOCX, TXT or MD.")


def chunk_text(text: str) -> list[str]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
    )
    return [chunk for chunk in splitter.split_text(text) if chunk.strip()]


def ingest_document(
    document_id: str,
    title: str,
    category: str,
    allowed_roles: list[str],
    text: str,
) -> int:
    chunks = chunk_text(text)
    if not chunks:
        raise EmptyDocument(
            "The document has no readable text. If it is a scanned PDF, upload a version with selectable text."
        )

    # Prefixing the title gives each chunk context about where it came from.
    vectors = embed_passages([f"{title}\n\n{chunk}" for chunk in chunks])

    vectorstore.ensure_indexes()
    vectorstore.replace_document_chunks(
        document_id,
        [
            {
                "document_id": document_id,
                "chunk_index": index,
                "title": title,
                "category": category,
                "allowed_roles": allowed_roles,
                "text": chunk,
                "embedding": vector,
            }
            for index, (chunk, vector) in enumerate(zip(chunks, vectors))
        ],
    )
    return len(chunks)
